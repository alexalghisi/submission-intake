"""Turns attachments into text and rows, and decides what kind of document each one is."""

from __future__ import annotations

import csv
import io
import re

import openpyxl
import pdfplumber

from .email_parser import Attachment
from .extract.spreadsheets import LOSS_COLUMNS, SOV_COLUMNS, find_header
from .models import DocKind, Document

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_attachment(doc_id: str, attachment: Attachment) -> Document:
    name = attachment.filename.lower()
    text, rows = "", []
    if name.endswith(".pdf") or attachment.content_type == "application/pdf":
        with pdfplumber.open(io.BytesIO(attachment.data)) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    elif name.endswith(".xlsx") or attachment.content_type == XLSX:
        workbook = openpyxl.load_workbook(io.BytesIO(attachment.data), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        rows = [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        text = "\n".join("\t".join(r) for r in rows)
    elif name.endswith(".csv") or attachment.content_type == "text/csv":
        decoded = attachment.data.decode("utf-8-sig", errors="replace")
        rows = [r for r in csv.reader(io.StringIO(decoded)) if any(c.strip() for c in r)]
        text = decoded
    elif attachment.content_type.startswith("text/"):
        text = attachment.data.decode("utf-8", errors="replace")
    document = Document(
        id=doc_id, filename=attachment.filename, content_type=attachment.content_type, text=text, rows=rows
    )
    return document.model_copy(update={"kind": classify(document)})


APPLICATION_MARKERS = re.compile(
    r"(commercial insurance application|applicant information|named insured|proposed eff)", re.I
)


def classify(document: Document) -> DocKind:
    """Rules first: they are cheap, explainable and right for the common formats."""
    # The same header detection the parsers use, so a sheet is only called an SOV if the SOV parser can read it.
    if document.rows:
        if find_header(document.rows, LOSS_COLUMNS) is not None:
            return DocKind.LOSS_RUN
        if find_header(document.rows, SOV_COLUMNS) is not None:
            return DocKind.SOV
    if APPLICATION_MARKERS.search(document.text):
        return DocKind.APPLICATION
    if re.search(r"loss run|loss history", document.filename + " " + document.text[:500], re.I):
        return DocKind.LOSS_RUN
    return DocKind.OTHER
